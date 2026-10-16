#!/usr/bin/env python3
"""Gemeindeverzeichnis von Destatis (GV-ISys, Gemeinden mit Fläche, Bevölkerung, Postleitzahl und
geografischen Mittelpunktkoordinaten) -> research/gemeinden.json.

Quelle: Statistisches Bundesamt, „Alle politisch selbständigen Gemeinden mit ausgewählten Merkmalen“,
Datenlizenz Deutschland – Namensnennung – Version 2.0. Rejouable : python3 scripts/data/fetch-gemeinden.py
"""
import json, os, sys, urllib.request, datetime
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
URL = "https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/Administrativ/Archiv/GVAuszugQ/AuszugGV2QAktuell.xlsx?__blob=publicationFile"
raw = os.path.join(ROOT, "research", "gemeinden-destatis.xlsx")
if not os.path.exists(raw) or "--refresh" in sys.argv:
    req = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0 (solarrechnen.de data script)"})
    open(raw, "wb").write(urllib.request.urlopen(req, timeout=120).read())
wb = openpyxl.load_workbook(raw, read_only=True)
ws = wb[wb.sheetnames[1]]
title = None
out = []
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: title = r[0]
    if r[0] != "60": continue
    plz, lon, lat = r[13], r[14], r[15]
    if not plz or not lon: continue
    out.append({"ars": f"{r[2]}{r[3]}{r[4]}{r[5]}{r[6]}", "name": r[7], "land": r[2], "plz": str(plz).zfill(5),
                "lat": float(str(lat).replace(",", ".")), "lon": float(str(lon).replace(",", ".")), "ew": int(r[9] or 0)})
json.dump({"source": URL, "title": title, "retrieved_at": datetime.date.today().isoformat(), "gemeinden": out},
          open(os.path.join(ROOT, "research", "gemeinden.json"), "w"), ensure_ascii=False)
print(len(out), "Gemeinden", title)
